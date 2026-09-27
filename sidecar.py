"""Python sidecar for KonkurentApp (Flutter Desktop).

Bridges the Flutter UI and the existing recognition engine in recog.py /
normalizer.py. Communication is one JSON object per invocation: result on
stdout, progress/errors folded into the JSON (never raw on stdout).

Usage:
    python sidecar.py --action=recognize --file=/abs/path/to/invoice.pdf
    python sidecar.py --action=export --output=/abs/path/to/out.xlsx   < state.json
    python sidecar.py --action=import --file=/abs/path/to/saved.xlsx
"""

import os
import sys
import json
import argparse
import contextlib

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

# Ссылка на исходный stdout: результат должен попасть сюда даже внутри
# redirect_stdout (recog печатает в stderr, но _fail может вызываться из
# обёрнутого блока).
_REAL_STDOUT = sys.stdout

META_LABEL_TO_KEY = {
    "Договор:": "contract",
    "Сроки поставки:": "deliveryTime",
    "Доставка:": "delivery",
    "Условия оплаты:": "paymentTerms",
    "Комментарий:": "comment",
}
META_KEY_TO_LABEL = {v: k for k, v in META_LABEL_TO_KEY.items()}


class _Table:
    """Mimics the camelot table interface expected by normalizer."""

    def __init__(self, df):
        self.df = df


def _configure_stdio():
    # On Windows the default stdio encoding is the ANSI code page; the Flutter
    # bridge expects UTF-8 on both stdout (JSON result) and stderr (recog/camelot
    # logs). Force all three streams to UTF-8.
    for stream in (sys.stdout, sys.stderr, sys.stdin):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass


def _emit(obj):
    try:
        text = json.dumps(obj, ensure_ascii=False)
        _REAL_STDOUT.write(text + "\n")
    except (UnicodeEncodeError, ValueError):
        _REAL_STDOUT.write(json.dumps(obj, ensure_ascii=True) + "\n")
    _REAL_STDOUT.flush()


@contextlib.contextmanager
def _stdout_to_stderr():
    """Keep stdout reserved for the single JSON result; recog's prints go to stderr."""
    with contextlib.redirect_stdout(sys.stderr):
        yield


def _fail(message):
    _emit({"status": "error", "message": message})
    sys.exit(1)


def _cell_str(value):
    if value is None:
        return ""
    return str(value).strip()


def _num_or_none(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    if not s or s.lower() in ("nan", "none", "-"):
        return None
    try:
        return float(s)
    except (ValueError, TypeError):
        return None


def _fmt_num(value):
    if value is None:
        return ""
    return value


def load_config():
    with open(os.path.join(SCRIPT_DIR, "config.json"), "r", encoding="utf-8") as f:
        return json.load(f)


def _tables_for(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        import recog
        return recog.extract_with_camelot(path)

    if ext in (".xlsx", ".xls", ".xlsm"):
        import pandas as pd
        sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
        tables = []
        for _, df in sheets.items():
            tables.append(_Table(df.fillna("")))
        return tables

    if ext == ".docx":
        import pandas as pd
        try:
            from docx import Document
        except ImportError:
            _fail("Для распознавания DOCX установите python-docx")
        doc = Document(path)
        tables = []
        for t in doc.tables:
            rows = [[cell.text for cell in row.cells] for row in t.rows]
            if rows:
                tables.append(_Table(pd.DataFrame(rows)))
        return tables

    _fail(f"Неподдерживаемый формат: {ext or 'без расширения'}")


def action_recognize(path):
    if not os.path.isfile(path):
        _fail(f"Файл не найден: {path}")

    from normalizer import process_pdf_tables, STANDARD_COLUMNS

    filename = os.path.basename(path)
    with _stdout_to_stderr():
        tables = _tables_for(path)
    if not tables:
        _fail("Таблицы в файле не найдены")

    df, _orig, bez_nds = process_pdf_tables(tables, filename)
    if df.shape[0] == 0:
        _fail("Не удалось распознать позиции в файле")

    offers = []
    for i in range(len(df)):
        row = {col: df.iloc[i][(filename, col)] for col in STANDARD_COLUMNS}
        offers.append({
            "itemName": _cell_str(row.get("Товар")),
            "qty": _num_or_none(row.get("Кол-во")),
            "unit": _cell_str(row.get("Ед. изм")),
            "sumWithVat": _num_or_none(row.get("Сумма")),
            "sumWithoutVat": _num_or_none(bez_nds[i] if i < len(bez_nds) else None),
        })

    _emit({
        "status": "ok",
        "supplierName": os.path.splitext(filename)[0],
        "sourceFileName": filename,
        "offers": offers,
    })


def action_export(output):
    raw_bytes = sys.stdin.buffer.read() if hasattr(sys.stdin, "buffer") else b""
    try:
        raw = raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        raw = raw_bytes.decode("utf-8", "replace")
    if not raw.strip():
        _fail("Пустое состояние на stdin")
    try:
        state = json.loads(raw)
    except json.JSONDecodeError as e:
        _fail(f"Некорректный JSON состояния: {e}")

    import pandas as pd
    import recog
    from normalizer import STANDARD_COLUMNS

    request_name = (state.get("requestName") or "").strip() or "Заявка"
    request_items = [
        (str(it.get("name", "")), _fmt_num(it.get("qty")))
        for it in state.get("requestItems", [])
        if str(it.get("name", "")).strip()
    ]

    pdf_data_list = []
    block_names = {}
    meta_values = {}

    for idx, sup in enumerate(state.get("suppliers", [])):
        display = (sup.get("displayName") or "").strip()
        source = (sup.get("sourceFileName") or "").strip()
        filename = source or f"{display or 'supplier'}.xlsx"
        if filename in block_names:
            filename = f"{idx}_{filename}"

        rows = []
        bez_nds = []
        for i, off in enumerate(sup.get("offers", [])):
            rows.append([
                str(i + 1),
                off.get("itemName") or "",
                _fmt_num(off.get("qty")),
                off.get("unit") or "",
                "",
                _fmt_num(off.get("sumWithVat")),
            ])
            bez_nds.append(_fmt_num(off.get("sumWithoutVat")))

        df = pd.DataFrame(
            rows,
            columns=pd.MultiIndex.from_product([[filename], STANDARD_COLUMNS]),
        )
        orig = {col: col for col in STANDARD_COLUMNS}
        pdf_data_list.append((df, orig, bez_nds))

        block_names[filename] = display or filename
        meta = sup.get("meta") or {}
        meta_values[filename] = {
            META_KEY_TO_LABEL[key]: meta.get(key, "")
            for key in META_LABEL_TO_KEY.values()
        }

    output = os.path.abspath(output)
    os.makedirs(os.path.dirname(output), exist_ok=True)

    with _stdout_to_stderr():
        recog.fill_template(
            pdf_data_list,
            os.path.dirname(output),
            SCRIPT_DIR,
            output_path=output,
            block_names=block_names,
            request_name=request_name,
            request_items=request_items or None,
            meta_values=meta_values,
        )

    _emit({"status": "ok", "path": output})


def action_import(path):
    if not os.path.isfile(path):
        _fail(f"Файл не найден: {path}")

    from openpyxl import load_workbook

    config = load_config()
    sheet_name = config["sheet_name"]

    wb = load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        _emit({"status": "foreign"})
        return
    ws = wb[sheet_name]

    a1 = ws.cell(row=1, column=1).value
    if not (isinstance(a1, str) and a1.strip().startswith("Заявка:")):
        _emit({"status": "foreign"})
        return

    request_name = a1.split(":", 1)[1].strip()

    n_fixed = len(config["fixed_columns"])
    block_size = len(config["block_columns"])
    meta_labels = [m["label"] for m in config["row"]["meta"]]
    fixed_header = config["fixed_columns"][0]["header"]

    header_row = None
    for r in range(1, ws.max_row + 1):
        if _cell_str(ws.cell(row=r, column=1).value) == fixed_header:
            header_row = r
            break
    if header_row is None:
        _fail("Не найдена строка заголовков в файле")

    data_start = header_row + 1
    hidden_row = header_row - 1 if header_row >= 3 else None

    meta_start = None
    if meta_labels:
        for r in range(data_start, ws.max_row + 1):
            if _cell_str(ws.cell(row=r, column=1).value) == meta_labels[0]:
                meta_start = r
                break
    if meta_start is None:
        meta_start = data_start

    blocks = []
    col = n_fixed + 1
    while col <= ws.max_column:
        name = ws.cell(row=1, column=col).value
        if name in (None, ""):
            break
        blocks.append({"start": col, "name": _cell_str(name)})
        col += block_size

    suppliers = []
    for b in blocks:
        cs = b["start"]
        source = None
        if hidden_row:
            v = ws.cell(row=hidden_row, column=cs).value
            if v not in (None, ""):
                source = _cell_str(v)

        offers = []
        for r in range(data_start, meta_start):
            item = ws.cell(row=r, column=cs).value
            qty = ws.cell(row=r, column=cs + 1).value
            unit = ws.cell(row=r, column=cs + 2).value
            sum_wo = ws.cell(row=r, column=cs + 4).value
            sum_w = ws.cell(row=r, column=cs + 5).value
            if item in (None, "") and qty in (None, "") and sum_w in (None, ""):
                continue
            offers.append({
                "itemName": _cell_str(item),
                "qty": _num_or_none(qty),
                "unit": _cell_str(unit),
                "sumWithVat": _num_or_none(sum_w),
                "sumWithoutVat": _num_or_none(sum_wo),
            })

        meta = {}
        for i, label in enumerate(meta_labels):
            v = ws.cell(row=meta_start + i, column=cs).value
            key = META_LABEL_TO_KEY.get(label)
            if key:
                meta[key] = _cell_str(v)

        suppliers.append({
            "displayName": b["name"],
            "sourceFileName": source,
            "offers": offers,
            "meta": meta,
        })

    request_items = []
    for r in range(data_start, meta_start):
        name = ws.cell(row=r, column=2).value
        qty = ws.cell(row=r, column=3).value
        if name not in (None, ""):
            request_items.append({
                "name": _cell_str(name),
                "qty": _cell_str(qty),
            })

    _emit({
        "status": "ok",
        "requestName": request_name,
        "requestItems": request_items,
        "suppliers": suppliers,
    })


def main():
    _configure_stdio()
    parser = argparse.ArgumentParser(description="KonkurentApp Python sidecar")
    parser.add_argument("--action", required=True,
                        choices=["recognize", "export", "import"])
    parser.add_argument("--file")
    parser.add_argument("--output")
    args = parser.parse_args()

    try:
        if args.action == "recognize":
            if not args.file:
                _fail("--file обязателен для recognize")
            action_recognize(args.file)
        elif args.action == "export":
            if not args.output:
                _fail("--output обязателен для export")
            action_export(args.output)
        elif args.action == "import":
            if not args.file:
                _fail("--file обязателен для import")
            action_import(args.file)
    except SystemExit:
        raise
    except Exception as e:
        _fail(f"{type(e).__name__}: {e}")


if __name__ == "__main__":
    main()